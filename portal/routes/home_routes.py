import io

from flask import Blueprint, render_template, request, redirect, url_for, current_app, flash, send_file
from flask_login import login_required, current_user
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
import db
import fx
import parser

home_bp = Blueprint("home", __name__)

@home_bp.route("/")
@login_required
def index():
    conn = db.get_connection(current_app.config["DB_PATH"])
    active_version = db.get_active_bom_version(conn)
    display_version = db.get_display_bom_version(conn)
    display_bom_id = display_version["id"] if display_version else None
    vehicles = db.list_vehicles(conn, display_bom_id)
    selected_vehicle = request.args.get("vehicle", "NE2_NV1")
    if selected_vehicle not in vehicles:
        selected_vehicle = vehicles[0] if vehicles else ""
    usd = fx.get_latest_rate(conn, "USD")
    eur = fx.get_latest_rate(conn, "EUR")
    report = db.dashboard_report(conn, display_bom_id, selected_vehicle)
    for major in report["majors"]:
        dom_total = {"material": 0.0, "logistics": 0.0, "tariff": 0.0, "total": 0.0}
        over_total = {c: {"material_lp": 0.0, "material_kd": 0.0, "material_total": 0.0, "logistics": 0.0, "tariff": 0.0, "total": 0.0} for c in report["overseas_countries"]}
        for cat in major["categories"]:
            for key in dom_total:
                dom_total[key] += cat["domestic"][key]
            for country in report["overseas_countries"]:
                for key in over_total[country]:
                    over_total[country][key] += cat["overseas"][country][key]
        major["domestic_total"] = dom_total
        major["overseas_total"] = over_total
    latest_bom_upload = db.get_latest_bom_upload(conn)
    conn.close()
    return render_template(
        "dashboard_new.html",
        usd=usd, eur=eur,
        latest_bom_upload=latest_bom_upload,
        report=report,
        vehicles=vehicles,
        selected_vehicle=selected_vehicle,
        active_version=active_version,
    )

@home_bp.route("/export/dashboard.xlsx")
@login_required
def export_dashboard():
    conn = db.get_connection(current_app.config["DB_PATH"])
    display_version = db.get_display_bom_version(conn)
    display_bom_id = display_version["id"] if display_version else None
    vehicles = db.list_vehicles(conn, display_bom_id)
    selected_vehicle = request.args.get("vehicle", "NE2_NV1")
    if selected_vehicle not in vehicles:
        selected_vehicle = vehicles[0] if vehicles else ""
    report = db.dashboard_report(conn, display_bom_id, selected_vehicle)
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "재료비 Summary"
    overseas = report["overseas_countries"]
    FONT_NAME = "현대하모니 L"
    HEADER_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    THIN_BOTTOM = Border(bottom=Side(style="thin"))
    bold_center = Alignment(horizontal="center", vertical="center")
    header_font = Font(name=FONT_NAME, bold=True)
    data_font = Font(name=FONT_NAME)

    def style_header(cell):
        cell.font = header_font
        cell.fill = HEADER_FILL
        cell.alignment = bold_center
        cell.border = THIN_BOTTOM

    title = f"▶ {selected_vehicle} 공조/열관리 시스템 재료비 (국내" + "".join(f"/{c}" for c in overseas) + ")"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6 + len(overseas) * 6)
    ws.cell(row=1, column=1, value=title).font = Font(name=FONT_NAME, size=14, bold=True)

    header_row, sub_row = 3, 4
    ws.merge_cells(start_row=header_row, start_column=1, end_row=sub_row, end_column=1)
    style_header(ws.cell(row=header_row, column=1, value="대분류"))
    ws.merge_cells(start_row=header_row, start_column=2, end_row=sub_row, end_column=2)
    style_header(ws.cell(row=header_row, column=2, value="구분"))

    col = 3
    ws.merge_cells(start_row=header_row, start_column=col, end_row=header_row, end_column=col + 3)
    style_header(ws.cell(row=header_row, column=col, value="내수"))
    for label in ("재료비", "물류비", "관세", "합계"):
        style_header(ws.cell(row=sub_row, column=col, value=label))
        col += 1

    country_start_cols = {}
    for country in overseas:
        country_start_cols[country] = col
        ws.merge_cells(start_row=header_row, start_column=col, end_row=header_row, end_column=col + 5)
        style_header(ws.cell(row=header_row, column=col, value=country))
        ws.merge_cells(start_row=sub_row, start_column=col, end_row=sub_row, end_column=col + 2)
        style_header(ws.cell(row=sub_row, column=col, value="재료비(LP/KD)"))
        style_header(ws.cell(row=sub_row + 1, column=col, value="LP"))
        style_header(ws.cell(row=sub_row + 1, column=col + 1, value="KD"))
        style_header(ws.cell(row=sub_row + 1, column=col + 2, value="합계"))
        for offset, label in enumerate(("물류비", "관세", "합계"), start=3):
            style_header(ws.cell(row=sub_row, column=col + offset, value=label))
        col += 6

    row = sub_row + 2
    for major_idx, major in enumerate(report["majors"]):
        first_row = row
        band_fill = HEADER_FILL if major_idx % 2 == 0 else None
        for cat in major["categories"]:
            d = cat["domestic"]
            values = {2: cat["category"], 3: d["material"], 4: d["logistics"], 5: d["tariff"], 6: d["total"]}
            for country in overseas:
                o = cat["overseas"][country]
                base = country_start_cols[country]
                values.update({base: o["material_lp"], base + 1: o["material_kd"], base + 2: o["material_total"], base + 3: o["logistics"], base + 4: o["tariff"], base + 5: o["total"]})
            for column, value in values.items():
                cell = ws.cell(row=row, column=column, value=value)
                cell.font = data_font
                if column > 2:
                    cell.alignment = Alignment(horizontal="right")
                if band_fill:
                    cell.fill = band_fill
            row += 1
        if row - 1 > first_row:
            ws.merge_cells(start_row=first_row, start_column=1, end_row=row - 1, end_column=1)
        major_cell = ws.cell(row=first_row, column=1, value=major["major"])
        major_cell.font = header_font
        major_cell.alignment = bold_center
        if band_fill:
            major_cell.fill = band_fill

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    style_header(ws.cell(row=row, column=1, value="합 계"))
    gd = report["grand_domestic"]
    totals = {3: gd["material"], 4: gd["logistics"], 5: gd["tariff"], 6: gd["total"]}
    for country in overseas:
        go = report["grand_overseas"][country]
        base = country_start_cols[country]
        totals.update({base: go["material_lp"], base + 1: go["material_kd"], base + 2: go["material_total"], base + 3: go["logistics"], base + 4: go["tariff"], base + 5: go["total"]})
    for column, value in totals.items():
        style_header(ws.cell(row=row, column=column, value=value))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"재료비_Summary_{selected_vehicle}.xlsx"
    return send_file(buffer, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@home_bp.route("/upload", methods=["POST"])
@login_required
def upload():
    flash("BOM 업로드는 관리자 화면에서 새 초안을 만든 뒤 진행해 주세요.")
    return redirect(url_for("admin.bom_versions") if current_user.role == "admin" else url_for("bom.index"))
