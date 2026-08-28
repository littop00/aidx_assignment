import openpyxl
from columns import DESIGN_FIELDS, LEVEL_COLUMNS, FIRST_DATA_ROW, COUNTRIES, purchase_column
import db

def parse_and_upsert(conn, file_like, bom_id=None, replace_existing=False):
    wb = openpyxl.load_workbook(file_like, data_only=True)
    ws = wb["BOM_복사본"]

    if replace_existing and bom_id is not None:
        db.replace_bom_draft_data(conn, bom_id)
    inserted, updated, skipped = 0, 0, []
    for row_num in range(FIRST_DATA_ROW, ws.max_row + 1):
        part_no_cell = ws.cell(row=row_num, column=13).value
        part_no = str(part_no_cell).strip() if part_no_cell else ""
        # A BOM's structural/header rows commonly have a part name but no
        # part number.  They still define the BOM hierarchy and must remain
        # in the portal in the exact same order as the source workbook.
        part_name_cell = ws.cell(row=row_num, column=12).value
        part_name = str(part_name_cell).strip() if part_name_cell else ""
        if not part_name:
            skipped.append(row_num)
            continue

        fields = {}
        for name, col in DESIGN_FIELDS:
            value = ws.cell(row=row_num, column=col).value
            fields[name] = str(value) if value is not None else None

        level_depth, level_marker = None, None
        for depth, col in enumerate(LEVEL_COLUMNS):
            value = ws.cell(row=row_num, column=col).value
            if value is not None and str(value).strip() != "":
                level_depth, level_marker = depth, str(value)
                break

        existed = db.get_part(conn, part_no, row_num, bom_id) is not None
        db.upsert_part(conn, part_no, row_num, level_depth, level_marker, fields, bom_id)
        # Sourcing columns (AM/AN and each country block) are BOM-owned data.
        # Preserve them so MIP rows can be displayed but excluded from user input.
        for country in COUNTRIES:
            source = ws.cell(row=row_num, column=purchase_column(country, "sourcing_part")).value
            assembly = ws.cell(row=row_num, column=purchase_column(country, "sourcing_assembly")).value
            if source is not None or assembly is not None:
                db.upsert_purchase(conn, part_no, row_num, country, {
                    "sourcing_part": str(source).strip() if source is not None else None,
                    "sourcing_assembly": str(assembly).strip() if assembly is not None else None,
                }, bom_id=bom_id)
        if existed:
            updated += 1
        else:
            inserted += 1

    return {"inserted": inserted, "updated": updated, "skipped": skipped}
