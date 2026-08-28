import openpyxl
from copy import copy
from io import BytesIO
from columns import DESIGN_FIELDS, PURCHASE_FIELDS, COUNTRY_BASE_COL
import db

COUNTRY_BLOCK_WIDTH = 13
KOREA_BLOCK_START = COUNTRY_BASE_COL["한국"]
EXISTING_COUNTRY_BLOCKS = {"한국": 39, "미국": 52, "유럽": 65}
FIRST_DYNAMIC_BLOCK_COLUMN = 78

def _copy_country_block(ws, target_column, country):
    """Insert a country cost block by copying Korea's 13-column layout and formatting."""
    source_merges = [copy(merged) for merged in ws.merged_cells.ranges
                     if merged.min_col >= KOREA_BLOCK_START and merged.max_col < KOREA_BLOCK_START + COUNTRY_BLOCK_WIDTH]
    shifted_merges = [copy(merged) for merged in ws.merged_cells.ranges if merged.min_col >= target_column]
    for merged in shifted_merges:
        ws.unmerge_cells(str(merged))
    ws.insert_cols(target_column, COUNTRY_BLOCK_WIDTH)
    for merged in shifted_merges:
        ws.merge_cells(start_row=merged.min_row, start_column=merged.min_col + COUNTRY_BLOCK_WIDTH,
                       end_row=merged.max_row, end_column=merged.max_col + COUNTRY_BLOCK_WIDTH)
    max_rows = max(ws.max_row, 11)
    for offset in range(COUNTRY_BLOCK_WIDTH):
        source_column = KOREA_BLOCK_START + offset
        target = target_column + offset
        ws.column_dimensions[openpyxl.utils.get_column_letter(target)].width = ws.column_dimensions[openpyxl.utils.get_column_letter(source_column)].width
        for row in range(1, max_rows + 1):
            source = ws.cell(row=row, column=source_column)
            destination = ws.cell(row=row, column=target)
            if row <= 10 and source.value is not None:
                destination.value = source.value
            if source.has_style:
                destination._style = copy(source._style)
            destination.number_format = source.number_format
            destination.font = copy(source.font)
            destination.fill = copy(source.fill)
            destination.border = copy(source.border)
            destination.alignment = copy(source.alignment)
            destination.protection = copy(source.protection)
    for merged in source_merges:
        ws.merge_cells(start_row=merged.min_row, start_column=target_column + (merged.min_col - KOREA_BLOCK_START),
                       end_row=merged.max_row, end_column=target_column + (merged.max_col - KOREA_BLOCK_START))
    ws.cell(row=3, column=target_column, value=country)

def export_to_template(conn, template_path):
    wb = openpyxl.load_workbook(template_path)
    ws = wb["BOM_복사본"]

    overseas_countries = [country for country in db.list_bom_countries(conn) if country != "한국"]
    countries = ["한국"] + overseas_countries
    country_columns = dict(EXISTING_COUNTRY_BLOCKS)
    next_dynamic_column = FIRST_DYNAMIC_BLOCK_COLUMN
    for country in overseas_countries:
        if country in country_columns:
            continue
        _copy_country_block(ws, next_dynamic_column, country)
        country_columns[country] = next_dynamic_column
        next_dynamic_column += COUNTRY_BLOCK_WIDTH

    parts = db.list_parts(conn)
    # Manual portal rows do not exist in the source workbook. Insert a row at
    # the recorded BOM position and clone the next row's style before writing.
    for part in (p for p in parts if p.get("manual_row")):
        row_num = part["row_num"]
        ws.insert_rows(row_num, 1)
        for col in range(1, ws.max_column + 1):
            source = ws.cell(row=row_num + 1, column=col)
            target = ws.cell(row=row_num, column=col)
            if source.has_style:
                target._style = copy(source._style)
            target.number_format = source.number_format
            target.alignment = copy(source.alignment)
            target.border = copy(source.border)
            target.fill = copy(source.fill)
            target.font = copy(source.font)

    for part in parts:
        row_num = part["row_num"]
        for name, col in DESIGN_FIELDS:
            if name == "part_no":
                continue
            value = part.get(name)
            if value is not None:
                ws.cell(row=row_num, column=col, value=value)

        for country in countries:
            purchase = db.get_purchase(conn, part["part_no"], row_num, country)
            if not purchase:
                continue
            for name, _ in PURCHASE_FIELDS:
                value = purchase.get(name)
                if name in ("sourcing_part", "sourcing_assembly"):
                    location = purchase.get(f"{name}_location")
                    if value and location:
                        value = f"{value}({location})"
                if value is not None:
                    col = country_columns[country] + next(offset for field, offset in PURCHASE_FIELDS if field == name)
                    ws.cell(row=row_num, column=col, value=value)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
