import openpyxl
from io import BytesIO
import db
import exporter

def make_template_with_merge(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM_복사본"
    ws.merge_cells("B1:D1")
    ws["B1"] = "■ NE2_NV1 BOM"
    path = tmp_path / "template.xlsx"
    wb.save(path)
    return str(path)

def test_export_writes_part_values_at_original_coordinates(tmp_path):
    conn = db.get_connection(str(tmp_path / "test.db"))
    db.init_db(conn)
    db.upsert_part(conn, "P001", row_num=11, level_depth=1, level_marker="●",
                    fields={"part_name": "FILTER", "qty": "2"})
    db.upsert_purchase(conn, "P001", 11, "한국", {"currency": "KRW", "unit_price_material": "1000", "material_cost": "2000"})

    template_path = make_template_with_merge(tmp_path)
    output_bytes = exporter.export_to_template(conn, template_path)

    wb = openpyxl.load_workbook(BytesIO(output_bytes))
    ws = wb["BOM_복사본"]
    assert ws.cell(row=11, column=12).value == "FILTER"
    assert ws.cell(row=11, column=42).value == "1000"
    assert ws.cell(row=11, column=43).value == "2000"

def test_export_preserves_merged_cells(tmp_path):
    conn = db.get_connection(str(tmp_path / "test.db"))
    db.init_db(conn)
    template_path = make_template_with_merge(tmp_path)
    output_bytes = exporter.export_to_template(conn, template_path)
    wb = openpyxl.load_workbook(BytesIO(output_bytes))
    ws = wb["BOM_복사본"]
    assert "B1:D1" in [str(r) for r in ws.merged_cells.ranges]
    assert ws["B1"].value == "■ NE2_NV1 BOM"

def test_export_handles_duplicate_part_no_different_rows(tmp_path):
    conn = db.get_connection(str(tmp_path / "test.db"))
    db.init_db(conn)
    db.upsert_part(conn, "P001", row_num=11, level_depth=1, level_marker="●", fields={"part_name": "FILTER-A"})
    db.upsert_part(conn, "P001", row_num=13, level_depth=1, level_marker="●", fields={"part_name": "FILTER-B"})
    template_path = make_template_with_merge(tmp_path)
    output_bytes = exporter.export_to_template(conn, template_path)
    wb = openpyxl.load_workbook(BytesIO(output_bytes))
    ws = wb["BOM_복사본"]
    assert ws.cell(row=11, column=12).value == "FILTER-A"
    assert ws.cell(row=13, column=12).value == "FILTER-B"

def test_export_adds_a_cost_block_for_an_added_overseas_country(tmp_path):
    conn = db.get_connection(str(tmp_path / "test.db"))
    db.init_db(conn)
    db.upsert_part(conn, "P001", 11, 0, "●", {"part_name": "FILTER", "qty": "2"})
    db.upsert_purchase(conn, "P001", 11, "독일", {"currency": "EUR", "unit_price_material": "300", "material_cost": "600"})
    db.set_bom_countries(conn, 1, ["독일"])
    template_path = make_template_with_merge(tmp_path)

    output_bytes = exporter.export_to_template(conn, template_path)

    wb = openpyxl.load_workbook(BytesIO(output_bytes))
    ws = wb["BOM_복사본"]
    # First extra country block begins after the existing Korea/US/Europe layout.
    assert ws.cell(row=3, column=78).value == "독일"
    assert ws.cell(row=11, column=80).value == "EUR"
    assert ws.cell(row=11, column=81).value == "300"
    assert ws.cell(row=11, column=82).value == "600"

def test_export_copies_country_header_merges_and_shifts_following_headers(tmp_path):
    conn = db.get_connection(str(tmp_path / "test.db"))
    db.init_db(conn)
    db.upsert_part(conn, "P001", 11, 0, "●", {"part_name": "FILTER", "qty": "2"})
    db.upsert_purchase(conn, "P001", 11, "슬로바키아", {"currency": "EUR", "unit_price_material": "300"})
    db.set_bom_countries(conn, 1, ["슬로바키아"])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM_복사본"
    ws.merge_cells(start_row=3, start_column=39, end_row=3, end_column=51)
    ws.cell(row=3, column=39, value="한국")
    ws.merge_cells(start_row=3, start_column=78, end_row=3, end_column=80)
    ws.cell(row=3, column=78, value="후속 헤더")
    path = tmp_path / "template.xlsx"
    wb.save(path)

    output_bytes = exporter.export_to_template(conn, str(path))

    exported = openpyxl.load_workbook(BytesIO(output_bytes))["BOM_복사본"]
    merged = {str(item) for item in exported.merged_cells.ranges}
    assert "BZ3:CL3" in merged
    assert "CM3:CO3" in merged
    assert exported.cell(row=3, column=78).value == "슬로바키아"
    assert exported.cell(row=3, column=91).value == "후속 헤더"

def test_export_inserts_manual_portal_row(tmp_path):
    conn = db.get_connection(str(tmp_path / "test.db")); db.init_db(conn)
    db.upsert_part(conn, "P1", 11, 0, "●", {"part_name": "BASE"})
    db.add_manual_part(conn, 1, "P1", 11, "after", {"part_name": "ADDED", "part_no": "P2"})
    output = exporter.export_to_template(conn, make_template_with_merge(tmp_path))
    ws = openpyxl.load_workbook(BytesIO(output))["BOM_복사본"]
    assert ws.cell(12, 12).value == "ADDED"
