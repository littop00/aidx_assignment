import openpyxl
from io import BytesIO

def make_bom_fixture(rows):
    """rows: list of dict, e.g. {"part_no": "P001", "part_name": "FILTER", "qty": 2, "level_col": 6}"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM_복사본"
    r = 12
    for row in rows:
        ws.cell(row=r, column=2, value=row.get("vehicle", "NE2_NV1"))
        ws.cell(row=r, column=3, value=row.get("category", "HVAC"))
        if "level_col" in row:
            ws.cell(row=r, column=row["level_col"], value=row.get("level_marker", "●"))
        ws.cell(row=r, column=12, value=row.get("part_name"))
        ws.cell(row=r, column=13, value=row.get("part_no"))
        ws.cell(row=r, column=19, value=row.get("qty"))
        r += 1
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
